"""Chunked PROFECO ingestion. Column mapping is the single provider schema boundary."""
import hashlib
import os
import tempfile
import sqlite3
import json
import zipfile
from pathlib import Path
from datetime import datetime,date,timezone
from html.parser import HTMLParser
from urllib.parse import urljoin,urlparse
import requests
import pandas as pd
from tenacity import retry,stop_after_attempt,wait_exponential
from .categorizer import normalize

COLUMN_MAP={'PRODUCTO':'producto','PRESENTACION':'presentacion','MARCA':'marca','CATEGORIA':'categoria','PRECIO':'precio','FECHAREGISTRO':'fecha','CADENACOMERCIAL':'cadena','NOMBRECOMERCIAL':'tienda','ESTADO':'estado','MUNICIPIO':'municipio'}
REQUIRED={'producto','precio','fecha','estado'}
PORTAL='https://datos.profeco.gob.mx/datos_abiertos/'

def key_text(value): return normalize(str(value)).upper()
class Links(HTMLParser):
    def __init__(self): super().__init__(); self.links=[]; self.href=None; self.parts=[]
    def handle_starttag(self,tag,attrs):
        if tag=='a': self.href=dict(attrs).get('href'); self.parts=[]
    def handle_data(self,data):
        if self.href: self.parts.append(data)
    def handle_endtag(self,tag):
        if tag=='a' and self.href:
            self.links.append((self.href,' '.join(self.parts))); self.href=None

@retry(stop=stop_after_attempt(3),wait=wait_exponential(multiplier=1,max=4),reraise=True)
def discover(year=None):
    year=year or date.today().year
    response=requests.get(PORTAL,timeout=8)
    response.raise_for_status()
    parser=Links(); parser.feed(response.text)
    for candidate in (year,year-1):
        for href,title in parser.links:
            text=key_text(title)
            if 'QUIEN ES QUIEN EN LOS PRECIOS' in text and str(candidate) in text:
                url=urljoin(PORTAL,href)
                if urlparse(url).hostname!='datos.profeco.gob.mx': raise ValueError('Enlace de descarga fuera de PROFECO')
                return url
    raise ValueError('No se encontró el archivo anual vigente. Usa --file con una descarga oficial; no se inventa el token.')

@retry(stop=stop_after_attempt(3),wait=wait_exponential(multiplier=1,max=4),reraise=True)
def download(url,path):
    total=0
    with requests.get(url,stream=True,timeout=(8,60)) as response:
        response.raise_for_status()
        with open(path,'wb') as target:
            for chunk in response.iter_content(1024*1024):
                total+=len(chunk)
                if total>1024*1024*1024: raise ValueError('Archivo mayor a 1 GB; descarga un extracto local.')
                target.write(chunk)

def chunks(path,tempdir):
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as archive:
            if '[Content_Types].xml' in archive.namelist():
                # Excel is bounded for the prototype; large annual exports should use CSV.
                from openpyxl import load_workbook
                workbook=load_workbook(path,read_only=True,data_only=True)
                rows=workbook.active.iter_rows(values_only=True)
                columns=next(rows)
                buffer=[]
                for row in rows:
                    buffer.append(row)
                    if len(buffer)>=100000: yield pd.DataFrame(buffer,columns=columns); buffer=[]
                if buffer: yield pd.DataFrame(buffer,columns=columns)
                workbook.close()
                return
            member=next((m for m in archive.infolist() if m.filename.lower().endswith('.csv')),None)
            if not member or member.file_size>1024*1024*1024: raise ValueError('ZIP sin CSV válido o demasiado grande')
            extracted=Path(tempdir)/'source.csv'
            import shutil
            with archive.open(member) as source,extracted.open('wb') as target: shutil.copyfileobj(source,target)
            path=extracted
    # Probe the full stream without loading it to avoid committing partial UTF-8 data.
    encoding='utf-8-sig'
    try:
        with open(path,encoding=encoding) as source:
            while source.read(1024*1024): pass
    except UnicodeDecodeError: encoding='latin-1'
    yield from pd.read_csv(path,encoding=encoding,chunksize=100000,dtype=str)

def sync_prices(repository,file_path=None):
    meta_path='profecoMeta/estado'
    meta=repository.get(meta_path) or {}
    try:
        with tempfile.TemporaryDirectory(prefix='cochinito-profeco-') as tempdir:
            path=Path(file_path) if file_path else Path(tempdir)/'download'
            if not file_path: download(discover(),path)
            digest=hashlib.sha256()
            with path.open('rb') as source:
                for data in iter(lambda:source.read(1024*1024),b''): digest.update(data)
            fingerprint=digest.hexdigest()
            if meta.get('ultimaHuella')==fingerprint: return {'changed':False,'rows':0}
            staging=sqlite3.connect(Path(tempdir)/'staging.sqlite')
            staging.execute('CREATE TABLE prices (id TEXT PRIMARY KEY, day TEXT, value TEXT)')
            states={key_text(s) for s in os.getenv('PROFECO_ESTADOS','NUEVO LEON,TAMAULIPAS').split(',') if s.strip()}
            municipalities={key_text(s) for s in os.getenv('PROFECO_MUNICIPIOS','').split(',') if s.strip()}
            read=0
            try:
                for chunk in chunks(path,tempdir):
                    read+=len(chunk)
                    chunk.columns=[COLUMN_MAP.get(key_text(c).replace(' ',''),str(c)) for c in chunk.columns]
                    if not REQUIRED.issubset(chunk.columns) or not {'cadena','tienda'}.intersection(chunk.columns): raise ValueError('Faltan columnas obligatorias. Verifica COLUMN_MAP contra el diccionario oficial.')
                    chunk=chunk.fillna('')
                    for record in chunk.to_dict('records'):
                        if states and key_text(record['estado']) not in states: continue
                        if municipalities and key_text(record.get('municipio','')) not in municipalities: continue
                        store=record.get('tienda') or record.get('cadena')
                        if not store or not record['producto']: raise ValueError('Producto o tienda vacíos')
                        price=float(str(record['precio']).replace('$','').replace(',',''))
                        if not 0<price<1000000: raise ValueError('Precio inválido')
                        raw_day=str(record['fecha'])
                        day=pd.to_datetime(raw_day,dayfirst=not raw_day[:4].isdigit() or '-' not in raw_day,errors='raise').date().isoformat()
                        product=' '.join(str(record.get(k,'')) for k in ('producto','presentacion','marca')).strip()
                        key=hashlib.sha256((normalize(product)+'|'+normalize(store)+'|'+key_text(record.get('municipio',''))).encode()).hexdigest()
                        value={k:record.get(k,'') for k in COLUMN_MAP.values()}
                        value.update(producto=product,tienda=store,precio=price,fecha=day,fuente='profeco')
                        staging.execute('INSERT INTO prices VALUES (?,?,?) ON CONFLICT(id) DO UPDATE SET day=excluded.day,value=excluded.value WHERE excluded.day>prices.day',(key,day,json.dumps(value)))
                staging.commit()
                count=0
                cursor=staging.execute('SELECT id,value FROM prices')
                while rows:=cursor.fetchmany(400):
                    for key,value in rows:
                        # Source-specific IDs preserve newer household ticket prices.
                        repository.put('precios/profeco-'+key,json.loads(value)); count+=1
                meta.update(ultimaHuella=fingerprint,ultimaActualizacion=datetime.now(timezone.utc).isoformat(),filasLeidas=read,filasCargadas=count,errores=[])
                repository.put(meta_path,meta)
                return {'changed':True,'rows':count}
            finally: staging.close()
    except Exception as error:
        meta['errores']=[str(error)[:300]]
        repository.put(meta_path,meta)
        raise ValueError('No se actualizaron los precios de PROFECO. Revisa el archivo o usa --file; los precios anteriores siguen disponibles.') from error
